"""
Episode pipeline tests (offline, no Premiere, no whisper).

Run from the vault root:
    .venv/bin/python -m unittest discover -s 08_AUTOMATION/tests -v
"""

import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "Scripts"))

from docx import Document  # noqa: E402

from episode.ingest import ingest, split_sentences  # noqa: E402
from episode.align import align_sentences, normalize  # noqa: E402
from episode.timeline import build_timeline, _numbers, tc  # noqa: E402


def _words(text: str, start: float = 0.0, per_word: float = 0.4):
    """Fake whisper output: one token every `per_word` seconds."""
    out, t = [], start
    for tok in normalize(text):
        out.append({"token": tok, "start": t, "end": t + per_word * 0.9})
        t += per_word
    return out


def _docx(tmp: Path) -> Path:
    d = Document()
    d.add_paragraph("EPISODE TITLE", style="Title")
    d.add_paragraph("Audience: test").runs[0].italic = True
    d.add_heading("COLD OPEN", level=1)
    d.add_paragraph("In 2023, State Farm lost 14.1 billion dollars on insurance. That is a lot.")
    d.add_paragraph("ON-SCREEN GRAPHIC: State Farm loss $14.1B in 2023, source: statefarm.com").runs[0].italic = True
    d.add_heading("CHAPTER 1: THE BET", level=1)
    d.add_paragraph("ILLUSTRATION CARD (square, on screen)").runs[0].italic = True
    d.add_paragraph("Same contract, two directions").runs[0].italic = True
    d.add_paragraph("Picture a room full of people. Every one of them fears a fire.")
    d.add_heading("END", level=1)
    d.add_heading("Sources", level=1)
    d.add_paragraph("State Farm 2024 results: newsroom.statefarm.com")
    path = tmp / "script.docx"
    d.save(str(path))
    return path


class IngestTest(unittest.TestCase):
    def test_structure(self):
        with tempfile.TemporaryDirectory() as t:
            s = ingest(_docx(Path(t)))
        self.assertEqual(s["title"], "EPISODE TITLE")
        self.assertEqual([x["title"] for x in s["sections"]], ["COLD OPEN", "CHAPTER 1: THE BET"])
        narration = [b for b in s["blocks"] if b["kind"] == "narration"]
        self.assertEqual(len(narration), 4)
        cues = [b["cue"] for b in s["blocks"] if b["kind"] == "cue"]
        self.assertEqual(cues[0]["variant"], "data")
        self.assertEqual(cues[0]["source"], "statefarm.com")
        self.assertEqual(cues[1]["caption"], "Same contract, two directions")
        # back matter after END is reference, never narration
        self.assertEqual(s["references"], ["State Farm 2024 results: newsroom.statefarm.com"])

    def test_sentence_split_keeps_decimals_and_abbreviations(self):
        self.assertEqual(
            split_sentences("The U.S. lost 14.1 billion. Then it grew."),
            ["The U.S. lost 14.1 billion.", "Then it grew."])


class AlignTest(unittest.TestCase):
    def test_tolerates_voice_actor_drift(self):
        sents = [{"id": "B1", "text": "We'll end this video on that number."},
                 {"id": "B2", "text": "Here's where it began."}]
        # VO adds "but", which is not in the script
        words = _words("We'll end this video on that number, but here's where it began.")
        report = align_sentences(sents, words)
        self.assertEqual(report["well_matched"], 2)
        self.assertLess(sents[0]["end"], sents[1]["start"])

    def test_unmatched_sentence_is_interpolated(self):
        sents = [{"id": "B1", "text": "First sentence here."},
                 {"id": "B2", "text": "Fifty-one seconds."},
                 {"id": "B3", "text": "Third sentence here."}]
        words = _words("First sentence here.") + _words("51 seconds", 2.0) + _words("Third sentence here.", 4.0)
        align_sentences(sents, words)
        self.assertTrue(sents[0]["end"] <= sents[1]["start"] <= sents[2]["start"])


class TimelineTest(unittest.TestCase):
    def test_numbers_from_words_and_digits(self):
        self.assertTrue({"37", "19"} <= _numbers("thirty-seven out of a hundred, and 19%"))
        self.assertIn("14.1", _numbers("14.1 billion dollars"))

    def test_timecode(self):
        self.assertEqual(tc(61.5), "00:01:01:12")

    def test_data_chart_lands_on_the_sentence_reading_its_numbers(self):
        with tempfile.TemporaryDirectory() as t:
            s = ingest(_docx(Path(t)))
        sents = [dict(b) for b in s["blocks"] if b["kind"] == "narration"]
        words = _words(" ".join(x["text"] for x in sents))
        report = align_sentences(sents, words)
        tl = build_timeline(s, sents, {"file": "vo.wav", "duration": 60.0}, report)
        chart = next(g for g in tl["graphics"] if g["template"] == "chart")
        self.assertEqual(chart["sentence_idx"][0], 0)  # "In 2023, State Farm lost 14.1 billion..."
        self.assertEqual(chart["origin"], "client")
        # the stat sentence already has the client chart, so no duplicate suggestion
        self.assertFalse(any(g["template"] == "stat" for g in tl["graphics"]))
        self.assertTrue(any(g["template"] == "chapter" for g in tl["graphics"]))


if __name__ == "__main__":
    unittest.main()
